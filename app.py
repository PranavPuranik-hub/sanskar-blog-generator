from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, Response
from pydantic import BaseModel
from generate import generate_blog, fetch_drive_file_bytes
import traceback
import urllib.parse
import time
import os
import re
from collections import defaultdict
from typing import Optional, List
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import requests as http_requests

app = FastAPI(title="Sanskar NGO Blog Generator")
app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"]
)

# ── Rate limiter ──────────────────────────────────────────────────────────────
RATE_LIMIT_MAX    = 10
RATE_LIMIT_WINDOW = 60
_rate_data: dict[str, list[float]] = defaultdict(list)

def is_rate_limited(ip: str) -> bool:
    now = time.time()
    _rate_data[ip] = [t for t in _rate_data[ip] if now - t < RATE_LIMIT_WINDOW]
    if len(_rate_data[ip]) >= RATE_LIMIT_MAX:
        return True
    _rate_data[ip].append(now)
    return False
# ─────────────────────────────────────────────────────────────────────────────

class BlogRequest(BaseModel):
    text: str
    folder_url: Optional[str] = ""

@app.post("/generate")
def generate(request: BlogRequest, req_info: Request):
    client_ip = req_info.client.host if req_info.client else "unknown"
    if is_rate_limited(client_ip):
        return JSONResponse(status_code=429, content={
            "blog": "Error: Too many requests. Please wait a minute.",
            "images": [], "keywords": [], "total_photos": 0, "photo_count": 0
        })
    try:
        return generate_blog(request.text, request.folder_url or "")
    except Exception as e:
        traceback.print_exc()
        return JSONResponse(status_code=500, content={
            "blog": f"Error: {str(e)}", "images": [],
            "keywords": [], "total_photos": 0, "photo_count": 0
        })

@app.get("/photo-proxy")
def photo_proxy(url: str):
    decoded = urllib.parse.unquote(url)

    if decoded.startswith("drive:"):
        file_id = decoded[6:]
        if not file_id or "/" in file_id or "." in file_id.replace("-", "").replace("_", ""):
            return JSONResponse(status_code=403, content={"error": "Invalid file ID"})
        try:
            data, mime = fetch_drive_file_bytes(file_id)
            return Response(
                content=data,
                media_type=mime,
                headers={"Cache-Control": "private, max-age=3600"}
            )
        except Exception as e:
            return JSONResponse(status_code=500, content={"error": str(e)})

    if decoded.startswith("https://lh3.googleusercontent.com/"):
        try:
            r = http_requests.get(decoded, headers={
                "User-Agent": "Mozilla/5.0 Chrome/124.0.0.0",
                "Referer": "https://photos.google.com/"
            }, timeout=15, stream=True)
            content_type = r.headers.get("Content-Type", "image/jpeg")
            return StreamingResponse(
                r.iter_content(chunk_size=8192),
                media_type=content_type,
                headers={"Cache-Control": "private, max-age=3600"}
            )
        except Exception as e:
            return JSONResponse(status_code=500, content={"error": str(e)})

    return JSONResponse(status_code=403, content={"error": "Forbidden"})

@app.get("/", response_class=HTMLResponse)
def root():
    return HTML_PAGE

# ── DOCX export endpoint (pure Python) ──────────────────────────────────────

class PhotoRef(BaseModel):
    proxy_url: str
    alt: str

class DocxRequest(BaseModel):
    blog: str
    images: List[PhotoRef] = []

def _parse_blog_to_docx_elements(blog_text: str):
    lines = blog_text.split("\n")
    elements = []
    first = True
    title_done = False
    for raw in lines:
        t = raw.strip()
        if not t:
            continue
        if t == "[PHOTO]":
            elements.append(("photo", None))
            continue
        if first:
            first = False
            if "|" in t:
                parts = t.split("|", 1)
                elements.append(("title", parts[0].strip()))
                elements.append(("subtitle", parts[1].strip()))
            else:
                elements.append(("title", t))
            title_done = True
            continue
        if title_done and not any(e[0] == "subtitle" for e in elements) \
                and not t.startswith("*") and len(t) < 120 and "—" not in t:
            elements.append(("subtitle", t))
            title_done = False
            continue
        title_done = False
        if t.startswith("*") and t.endswith("*") and len(t) > 2:
            elements.append(("italic", t[1:-1]))
        elif t.startswith(">"):
            elements.append(("quote", t[1:].strip()))
        elif t.startswith("**") and t.endswith("**") and len(t) > 4:
            elements.append(("h2", t[2:-2]))
        elif t == t.upper() and re.search(r"[A-Z]{3}", t) and len(t) > 3 \
                and not re.match(r"^[·•\"'>*]", t) and "@" not in t:
            elements.append(("h1caps", t))
        elif re.match(r"^[·•]", t):
            elements.append(("bullet", re.sub(r"^[·•]\s*", "", t)))
        else:
            elements.append(("para", t))
    return elements

def _fetch_image_bytes(proxy_url: str):
    decoded = urllib.parse.unquote(proxy_url.replace("/photo-proxy?url=", ""))
    if decoded.startswith("drive:"):
        return fetch_drive_file_bytes(decoded[6:])
    else:
        r = http_requests.get(decoded, headers={"User-Agent": "Mozilla/5.0"}, timeout=15)
        r.raise_for_status()
        return r.content, r.headers.get("Content-Type", "image/jpeg")

@app.post("/export/docx")
def export_docx(req: DocxRequest):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    elements = _parse_blog_to_docx_elements(req.blog)
    photo_idx = 0
    total_photos = len(req.images)

    for elem_type, content in elements:
        if elem_type == "photo":
            if photo_idx < total_photos:
                img = req.images[photo_idx]
                try:
                    img_bytes, mime = _fetch_image_bytes(img.proxy_url)
                    ext = mime.split('/')[-1].lower()
                    if ext in ('jpeg', 'jpg', 'png'):
                        stream = BytesIO(img_bytes)
                        p = doc.add_paragraph()
                        p.add_run().add_picture(stream, width=Inches(5.5))
                        cap = doc.add_paragraph(f"📸 {img.alt}")
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.style.font.size = Pt(9)
                        cap.style.font.italic = True
                    else:
                        doc.add_paragraph(f"[Image format {ext} not supported: {img.alt}]")
                except Exception:
                    doc.add_paragraph(f"[Could not embed photo: {img.alt}]")
                photo_idx += 1
        elif elem_type == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(content)
            run.bold = True
            run.font.size = Pt(22)
            run.font.name = 'Georgia'
        elif elem_type == "subtitle":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(content)
            run.italic = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            p.paragraph_format.space_after = Pt(6)
            line = doc.add_paragraph()
            line_run = line.add_run("_" * 50)
            line_run.font.size = Pt(6)
            line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif elem_type == "italic":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.italic = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x4A, 0x4A, 0x6A)
        elif elem_type == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(content)
            run.italic = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x5A, 0x2A, 0x27)
        elif elem_type == "h1caps":
            p = doc.add_paragraph()
            run = p.add_run(content.upper())
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
            p.paragraph_format.space_before = Pt(12)
        elif elem_type == "h2":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
            p.paragraph_format.space_before = Pt(12)
        elif elem_type == "bullet":
            doc.add_paragraph(content, style='List Bullet')
        else:
            p = doc.add_paragraph(content)
            p.paragraph_format.space_after = Pt(6)

    while photo_idx < total_photos:
        img = req.images[photo_idx]
        try:
            img_bytes, mime = _fetch_image_bytes(img.proxy_url)
            ext = mime.split('/')[-1].lower()
            if ext in ('jpeg', 'jpg', 'png'):
                stream = BytesIO(img_bytes)
                p = doc.add_paragraph()
                p.add_run().add_picture(stream, width=Inches(5.5))
                cap = doc.add_paragraph(f"📸 {img.alt}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.style.font.size = Pt(9)
                cap.style.font.italic = True
            else:
                doc.add_paragraph(f"[Image format {ext} not supported: {img.alt}]")
        except Exception:
            doc.add_paragraph(f"[Could not embed photo: {img.alt}]")
        photo_idx += 1

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=sanskar-blog.docx"}
    )

# ── Frontend HTML (PDF button removed) ──────────────────────────────────────
HTML_PAGE = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Sanskar Blog Generator</title>
<link href="https://fonts.googleapis.com/css2?family=Playfair+Display:wght@700;800&family=DM+Sans:wght@400;500;600&display=swap" rel="stylesheet">
<style>
*{box-sizing:border-box;margin:0;padding:0}
:root{
  --ink:#1a1a2e;--muted:#6b7280;--accent:#c0392b;--accent2:#e67e22;
  --bg:#faf9f7;--card:#ffffff;--border:#e8e3dd;--success:#27ae60;
  --font-head:'Playfair Display',Georgia,serif;
  --font-body:'DM Sans',system-ui,sans-serif;
}
body{font-family:var(--font-body);background:var(--bg);color:var(--ink);min-height:100vh;padding:32px 16px}
.wrap{max-width:900px;margin:0 auto}

.site-header{text-align:center;margin-bottom:36px}
.site-header h1{font-family:var(--font-head);font-size:clamp(26px,4vw,38px);color:var(--ink);letter-spacing:-0.5px;line-height:1.2}
.site-header h1 span{color:var(--accent)}
.site-header p{color:var(--muted);font-size:14px;margin-top:6px}

.card{background:var(--card);border:1px solid var(--border);border-radius:16px;padding:24px;margin-bottom:20px;box-shadow:0 2px 8px rgba(0,0,0,.04)}
.card-label{font-size:11px;font-weight:600;text-transform:uppercase;letter-spacing:.8px;color:var(--muted);margin-bottom:10px;display:flex;align-items:center;gap:6px}
.card-label::before{content:'';display:inline-block;width:3px;height:14px;background:var(--accent);border-radius:2px}

textarea, input[type=text]{
  width:100%;border:1.5px solid var(--border);border-radius:10px;
  padding:12px 14px;font-size:14px;font-family:var(--font-body);color:var(--ink);
  background:#fdfcfb;resize:vertical;line-height:1.6;transition:border-color .2s
}
textarea{height:200px}
textarea:focus, input[type=text]:focus{outline:none;border-color:var(--accent)}
.input-hint{font-size:12px;color:var(--muted);margin-top:6px}

.drive-section{margin-top:16px;padding-top:16px;border-top:1px solid var(--border)}
.drive-label{font-size:12px;font-weight:600;color:var(--ink);margin-bottom:6px;display:flex;align-items:center;gap:6px}
.drive-label svg{flex-shrink:0}
.drive-note{background:#fff8f0;border:1px solid #fde8cc;border-radius:8px;padding:10px 14px;font-size:12px;color:#92400e;margin-top:8px;line-height:1.6}
.drive-note strong{display:block;margin-bottom:2px}

.btn-generate{
  width:100%;margin-top:16px;padding:14px;
  background:var(--accent);color:#fff;border:none;border-radius:10px;
  font-size:15px;font-weight:600;font-family:var(--font-body);
  cursor:pointer;transition:opacity .2s,transform .1s;letter-spacing:.2px
}
.btn-generate:hover:not(:disabled){opacity:.92;transform:translateY(-1px)}
.btn-generate:disabled{opacity:.55;cursor:not-allowed;transform:none}

#status{text-align:center;font-size:13px;color:var(--accent);margin-top:10px;min-height:18px;font-style:italic}
#out-card{display:none}

.blog{font-family:'Playfair Display',Georgia,serif;font-size:16px;line-height:1.95;color:#1a1a2e}
.b-title{font-size:clamp(22px,3vw,30px);font-weight:800;color:#1a1a2e;margin-bottom:6px;line-height:1.25;font-family:'Playfair Display',Georgia,serif}
.b-sub{font-size:14px;color:var(--muted);padding-bottom:16px;border-bottom:2px solid var(--border);margin-bottom:22px;font-family:'DM Sans',sans-serif;font-style:italic}
.b-para{margin-bottom:14px;font-family:'DM Sans',sans-serif;font-size:15px}
.b-italic{font-style:italic;color:#4a4a6a;font-size:17px;margin-bottom:16px;padding-left:2px;font-family:'Playfair Display',Georgia,serif}
.b-h-caps{font-size:11px;font-weight:700;color:var(--accent);text-transform:uppercase;letter-spacing:1.2px;margin:30px 0 12px;font-family:'DM Sans',sans-serif;padding-bottom:4px;border-bottom:2px solid var(--accent)}
.b-h-bold{font-size:19px;font-weight:700;color:#1a1a2e;margin:28px 0 12px;font-family:'Playfair Display',Georgia,serif}
.b-bul{display:flex;gap:10px;margin-bottom:9px;padding-left:4px;font-family:'DM Sans',sans-serif;font-size:15px}
.b-bul .dot{color:var(--accent);font-weight:900;flex-shrink:0;margin-top:2px}
.b-hl{margin-bottom:10px;padding-left:16px;border-left:3px solid var(--accent2);line-height:1.8;font-family:'DM Sans',sans-serif;font-size:15px}
.b-hl strong{color:var(--ink)}
.b-quote{border-left:4px solid var(--accent);padding:14px 20px;margin:22px 0;background:#fef9f9;border-radius:0 10px 10px 0;color:#5a2a27;font-style:italic;line-height:1.85;font-family:'Playfair Display',Georgia,serif;font-size:16px}

.photo{width:100%;margin:26px 0;border-radius:12px;overflow:hidden;border:1px solid var(--border);box-shadow:0 4px 16px rgba(0,0,0,.07)}
.photo img{width:100%;max-height:480px;object-fit:cover;display:block}
.photo-cap{font-size:12px;color:var(--muted);padding:8px 16px;font-style:italic;background:#faf9f7;font-family:'DM Sans',sans-serif}
.no-photo{height:90px;display:flex;align-items:center;justify-content:center;color:#c0b8b0;font-size:13px;border:2px dashed var(--border);border-radius:10px;margin:20px 0;font-family:'DM Sans',sans-serif}

.info-bar{background:#f0fdf4;border:1px solid #bbf7d0;border-radius:8px;padding:10px 14px;font-size:12px;color:#166534;margin-top:14px;font-family:'DM Sans',sans-serif}

.btn-row{display:flex;gap:8px;margin-top:14px;flex-wrap:wrap}
.dl-btn{
  flex:1;min-width:100px;padding:10px 12px;border:none;border-radius:8px;
  font-size:13px;font-weight:600;cursor:pointer;font-family:'DM Sans',sans-serif;
  transition:opacity .15s,transform .1s;white-space:nowrap
}
.dl-btn:hover:not(:disabled){opacity:.88;transform:translateY(-1px)}
.dl-btn:disabled{opacity:.55;cursor:not-allowed}
.btn-txt{background:#1a1a2e;color:#fff}
.btn-html{background:#2563eb;color:#fff}
.btn-docx{background:#16a34a;color:#fff}
.btn-copy{background:#8b5cf6;color:#fff}
</style>
</head>
<body>
<div class="wrap">

<div class="site-header">
  <h1>Sanskar <span>Blog</span> Generator</h1>
  <p>Paste event details · AI writes · Export in any format</p>
</div>

<div class="card">
  <div class="card-label">Event Details</div>
  <textarea id="inp" placeholder="Paste your event details here..."></textarea>

  <div class="drive-section">
    <div class="drive-label">
      <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#c0392b" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M22 19a2 2 0 0 1-2 2H4a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h5l2 3h9a2 2 0 0 1 2 2z"/></svg>
      Google Drive Folder URL <span style="color:var(--muted);font-weight:400">(optional — for photos)</span>
    </div>
    <input type="text" id="folder-url" placeholder="https://drive.google.com/drive/folders/YOUR_FOLDER_ID">
    <div class="drive-note">
      <strong>📋 How to connect your Drive folder:</strong>
      1. Share the Drive folder with your service account email (found in service_account.json → "client_email")<br>
      2. Give it "Viewer" access<br>
      3. Paste the folder link above — photos will be fetched securely via service account
    </div>
  </div>

  <button class="btn-generate" id="btn" onclick="go()">✨ Generate Blog</button>
  <div id="status"></div>
</div>

<div class="card" id="out-card">
  <div class="card-label">Blog — Ready to Publish</div>
  <div class="blog" id="blog"></div>
  <div class="info-bar" id="info"></div>
  <div class="btn-row">
    <button class="dl-btn btn-txt"  onclick="dlTxt()">📄 Plain Text</button>
    <button class="dl-btn btn-html" onclick="dlHTML()">🌐 HTML + Photos</button>
    <button class="dl-btn btn-docx" id="btn-docx" onclick="dlDOCX()">📝 Word (.docx)</button>
    <button class="dl-btn btn-copy" onclick="copyRichText()">📋 Copy Rich Text</button>
  </div>
</div>

</div>

<script>
let rawBlog = "", photos = [], pi = 0;

function nextPhoto(){
  if(pi < photos.length){
    const p = photos[pi++];
    return `<div class="photo">
      <img src="${p.url}" alt="${p.alt}"
           onerror="this.src='${p.thumb}';this.onerror=()=>this.parentElement.innerHTML='<div class=no-photo>📷 Photo unavailable</div>'">
      <div class="photo-cap">📸 ${p.alt}</div>
    </div>`;
  }
  return '';
}

function render(text){
  pi = 0;
  let html='', first=true, titleDone=false;
  for(const raw of text.split('\n')){
    const t = raw.trim();
    if(!t) continue;
    if(t==='[PHOTO]'){html+=nextPhoto();continue;}
    if(first){
      first=false;
      if(t.includes('|')){
        const parts=t.split('|');
        html+=`<div class="b-title">${parts[0].trim()}</div>`;
        if(parts[1]) html+=`<div class="b-sub">${parts.slice(1).join('|').trim()}</div>`;
      } else { html+=`<div class="b-title">${t}</div>`; }
      titleDone=true; continue;
    }
    if(titleDone && !html.includes('b-sub') && !t.startsWith('*') && t.length<120 && !t.includes('—')){
      html+=`<div class="b-sub">${t}</div>`; titleDone=false; continue;
    }
    titleDone=false;
    if(t.startsWith('*')&&t.endsWith('*')&&t.length>2){html+=`<div class="b-italic">${t.slice(1,-1)}</div>`;continue;}
    if(t.startsWith('>')){html+=`<div class="b-quote">${t.slice(1).trim()}</div>`;continue;}
    if(t.startsWith('**')&&t.endsWith('**')&&t.length>4){html+=`<div class="b-h-bold">${t.slice(2,-2)}</div>`;continue;}
    if(t===t.toUpperCase()&&/[A-Z]{3}/.test(t)&&t.length>3&&!/^[·•"'>*]/.test(t)&&!t.includes('@')){
      html+=`<div class="b-h-caps">${t}</div>`;continue;
    }
    if(/^[·•]/.test(t)){
      html+=`<div class="b-bul"><span class="dot">·</span><span>${t.replace(/^[·•]\s*/,'')}</span></div>`;continue;
    }
    if(t.includes(' — ')&&!t.startsWith('Join')&&!t.startsWith('Contact')){
      const idx=t.indexOf(' — ');
      html+=`<div class="b-hl"><strong>${t.substring(0,idx)}</strong> — ${t.substring(idx+3)}</div>`;continue;
    }
    html+=`<div class="b-para">${t}</div>`;
  }
  if(!text.includes('[PHOTO]')&&photos.length){
    photos.forEach(p=>{
      html+=`<div class="photo"><img src="${p.url}" alt="${p.alt}"
        onerror="this.src='${p.thumb}';this.onerror=()=>this.parentElement.innerHTML='<div class=no-photo>📷 Photo unavailable</div>'">
        <div class="photo-cap">📸 ${p.alt}</div></div>`;
    });
  }
  document.getElementById('blog').innerHTML=html;
}

async function go(){
  const text = document.getElementById('inp').value.trim();
  if(!text){alert('Paste event details first!');return;}
  const folderUrl = document.getElementById('folder-url').value.trim();
  const btn=document.getElementById('btn'), status=document.getElementById('status');
  btn.disabled=true; btn.textContent='⏳ Writing your story…';
  status.textContent='AI is generating your blog…';
  document.getElementById('out-card').style.display='none';
  try{
    const res=await fetch('/generate',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({text, folder_url: folderUrl})});
    if(res.status===429){status.textContent='⚠️ Too many requests — wait a minute.';return;}
    const data=await res.json();
    if(!data.blog||data.blog.startsWith('Error')){status.textContent='⚠️ '+(data.blog||'Unknown error');return;}
    rawBlog=data.blog; photos=Array.isArray(data.images)?data.images:[];
    render(rawBlog);
    document.getElementById('out-card').style.display='block';
    const wc=rawBlog.split(' ').length;
    const photoMsg=data.total_photos>0
      ?`📸 ${data.total_photos} photos in folder → ${data.photo_count} selected`
      :'⚠️ No photos — add a Drive folder URL above';
    document.getElementById('info').textContent=`✅ ${wc} words  •  ${photoMsg}`;
    status.textContent='✅ Done!';
    document.getElementById('out-card').scrollIntoView({behavior:'smooth'});
  }catch(e){status.textContent='❌ '+e.message;}
  finally{btn.disabled=false;btn.textContent='✨ Generate Blog';}
}

function dlTxt(){
  const a=document.createElement('a');
  a.href=URL.createObjectURL(new Blob([rawBlog],{type:'text/plain'}));
  a.download='sanskar-blog.txt'; a.click();
}

async function getBase64Images(){
  const map={};
  for(const p of photos){
    for(const src of [p.url,p.thumb]){
      if(map[src]) continue;
      try{
        const res=await fetch(src), blob=await res.blob();
        map[src]=await new Promise((ok,fail)=>{
          const fr=new FileReader();
          fr.onload=()=>ok(fr.result); fr.onerror=()=>fail();
          fr.readAsDataURL(blob);
        });
      }catch{map[src]=src;}
    }
  }
  return map;
}

async function dlHTML(){
  const status=document.getElementById('status');
  status.textContent='Embedding photos…';
  const imgMap=await getBase64Images();
  const clone=document.getElementById('blog').cloneNode(true);
  clone.querySelectorAll('img').forEach(img=>{
    const s=img.getAttribute('src');
    if(s&&imgMap[s]) img.setAttribute('src',imgMap[s]);
    const oe=img.getAttribute('onerror')||'';
    const m=oe.match(/this\.src='([^']+)'/);
    if(m&&imgMap[m[1]]) img.setAttribute('onerror',oe.replace(m[1],imgMap[m[1]]));
  });
  const titleEl=clone.querySelector('.b-title');
  const docTitle=titleEl?titleEl.textContent.trim():'Sanskar NGO Blog';
  const full=`<!DOCTYPE html><html><head><meta charset="UTF-8"><title>${docTitle}</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:'Georgia',serif;background:#fff;padding:40px;max-width:800px;margin:0 auto;line-height:1.8}
h1{font-size:28px;margin-bottom:8px}
.subtitle{font-size:14px;color:#6b7280;border-bottom:2px solid #eee;margin-bottom:20px;font-style:italic}
.photo{margin:30px 0;text-align:center}
.photo img{max-width:100%;border-radius:12px;box-shadow:0 2px 8px rgba(0,0,0,0.1)}
.photo-cap{font-size:12px;color:#6b7280;margin-top:6px}
blockquote{border-left:4px solid #c0392b;margin:20px 0;padding:10px 20px;background:#fef9f9;font-style:italic}
h2{font-size:20px;margin:24px 0 12px;color:#1a1a2e}
</style>
</head><body>
${clone.innerHTML}
<hr><p style="font-size:12px;color:#aaa;margin-top:40px">Generated by Sanskar NGO Blog Generator</p>
</body></html>`;
  const a=document.createElement('a');
  a.href=URL.createObjectURL(new Blob([full],{type:'text/html'}));
  a.download='sanskar-blog.html'; a.click();
  status.textContent='✅ HTML downloaded.';
}

async function dlDOCX(){
  const btn=document.getElementById('btn-docx');
  const status=document.getElementById('status');
  btn.disabled=true; btn.textContent='⏳ Building .docx…';
  status.textContent='Generating Word document…';
  try{
    const res=await fetch('/export/docx',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({blog:rawBlog, images: photos.map(p=>({proxy_url:p.url,alt:p.alt}))})
    });
    if(!res.ok) throw new Error(await res.text());
    const blob=await res.blob();
    const a=document.createElement('a');
    a.href=URL.createObjectURL(blob);
    a.download='sanskar-blog.docx'; a.click();
    status.textContent='✅ Word document downloaded!';
  }catch(e){
    status.textContent='❌ DOCX error: '+e.message;
  }finally{
    btn.disabled=false; btn.textContent='📝 Word (.docx)';
  }
}

async function copyRichText(){
  const status=document.getElementById('status');
  status.textContent='Copying rich text to clipboard…';
  const blogHtml = document.getElementById('blog').cloneNode(true);
  const imgMap = await getBase64Images();
  blogHtml.querySelectorAll('img').forEach(img=>{
    const src = img.getAttribute('src');
    if(src && imgMap[src]) img.setAttribute('src', imgMap[src]);
  });
  const htmlContent = blogHtml.outerHTML;
  const blob = new Blob([htmlContent], {type: 'text/html'});
  const clipboardItem = new ClipboardItem({
    'text/html': blob,
    'text/plain': new Blob([rawBlog], {type: 'text/plain'})
  });
  try {
    await navigator.clipboard.write([clipboardItem]);
    status.textContent = '✅ Copied! Paste into Google Docs, Word, or any rich text editor.';
  } catch (err) {
    await navigator.clipboard.writeText(rawBlog);
    status.textContent = '⚠️ Rich text copy failed, plain text copied instead.';
  }
}
</script>
</body>
</html>"""